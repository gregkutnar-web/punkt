#!/usr/bin/env python3
"""Zimski letaki Odiseje (LUNA, MARS, Balthazar): SVG + Excel + poročilo.

Vir vsebine: docx (prepis, popravljene le slovnične/tipkarske napake),
vir logotipov: SVG iz Illustratorja. Enote SVG so točke (pt), 1 pt = 1 uporabniška enota.

Uporaba: python3 build_letaki.py <Zimski_letak_bari.docx> <LOGO_VSI_LOKALI.svg> [izhodna_mapa]
"""
import math
import random
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", NS)

# ---- Format in CGP Odiseje -------------------------------------------------
MM = 72 / 25.4
W, H = round(148 * MM, 2), round(220 * MM, 2)  # 419.53 x 623.62 pt
GAP = round(20 * MM, 2)
M = 34.0  # rob besedila
NAVY, TEAL, GREEN, WHITE = "#071d22", "#00b1bc", "#a5cb39", "#ffffff"
SLOGAN = "Vesolje vsebin"
FOOT = [
    "Alergeni so navedeni v tabeli pri šanku. Cene so v evrih in veljajo od 1. 11. 2026. DDV je vključen v ceni.",
    "Odiseja d.o.o. Moskovska ulica 11, 1000 Ljubljana",
]
NBSP = " "

# Illustrator bere PostScript ime pisave; brskalnik ob nameščeni Archivo najde drugo družino v seznamu.
FONT = {
    400: ("Archivo-Regular", 400),
    500: ("Archivo-Medium", 500),
    600: ("Archivo-SemiBold", 600),
    700: ("Archivo-Bold", 700),
}

BARS = [
    # id, ime v docx, prikazno ime, logotip, poudarek naslovov, poudarek cen
    dict(id="luna", docx="LUNA", name="LUNA", logo="LUNA", head=TEAL, price=TEAL),
    dict(id="mars", docx="MARS", name="MARS", logo="MARS", head=GREEN, price=GREEN),
    dict(id="balthazar", docx="Baltazar", name="Balthazar", logo="BALTHAZAR", head=TEAL, price=GREEN),
]

# ---- Popravki besedila (samo slovnica in tipkarske napake) -----------------
FIXES = [
    (r"whiskey Screwball smetana", "whiskey Skrewball, smetana", "manjkajoča vejica; blagovna znamka Skrewball"),
    (r"chai late\b", "chai latte", "tipkarska napaka"),
    (r"ginger bread", "gingerbread", "tipkarska napaka (enotno z »gingerbread sirup«)"),
    (r"malininim", "malinovim", "slovnična napaka"),
    (r"Ramazoti rossato", "Ramazzotti Rosato", "tipkarska napaka"),
    (r"\bscrewball\b", "Skrewball", "blagovna znamka Skrewball, velika začetnica"),
    (r"\bScrewball\b", "Skrewball", "blagovna znamka Skrewball (enotno)"),
    (r"\bradenska\b", "Radenska", "blagovna znamka, velika začetnica"),
    (r"\s+,", ",", "odvečen presledek pred vejico"),
    (r",(?=[^\s\d)])", ", ", "manjkajoč presledek za vejico"),
    (r",\)", ")", "odvečna vejica pred zaklepajem"),
    (r"^\(Matcha,", "(matcha,", "velika začetnica v naštevanju"),
    (r"^\(Ube latte,", "(ube latte,", "velika začetnica v naštevanju"),
    (r"^\(Vroča čokolada,", "(vroča čokolada,", "velika začetnica v naštevanju"),
    (r"\s{2,}", " ", "dvojni presledek"),
]


def fix_text(s):
    why = []
    out = s.strip()
    if out.endswith(")") and not out.startswith("(") and "(" not in out:
        out = "(" + out
        why.append("manjkajoči oklepaj")
    for pat, rep, reason in FIXES:
        new = re.sub(pat, rep, out)
        if new != out:
            why.append(reason)
            out = new
    return out, why


# ---- Branje docx -----------------------------------------------------------
HEADINGS = {
    "Topli brezalkoholni napitki", "Alkoholni napitki", "Vroče čokolade",
    "Topli napitki", "Šprici", "Koktejli",
}
PRICE_RE = re.compile(r"^(.*?)\s*(\d+,\d)\s*€\s*$")


def read_docx(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf8")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S):
        t = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, flags=re.S))
        t = t.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        paras.append(t.replace(" ", " "))
    bars, bar, page, blanks = {}, None, None, 0
    names = {b["docx"].lower(): b["docx"] for b in BARS}
    for raw in paras:
        t = raw.strip()
        if not t:
            blanks += 1
            continue
        if t.lower() in names:
            bar = bars.setdefault(names[t.lower()], {"pages": []})
            page = None
            blanks = 0
            continue
        if t.lower() in ("prva stran", "druga stran"):
            page = {"sections": []}
            bar["pages"].append(page)
            blanks = 0
            continue
        if t in HEADINGS:
            page["sections"].append({"title": t, "items": []})
            blanks = 0
            continue
        m = PRICE_RE.match(raw.strip())
        if m:
            if not page["sections"]:
                page["sections"].append({"title": None, "items": []})
            sec = page["sections"][-1]
            gap = blanks >= 2 and bool(sec["items"])
            sec["items"].append({"name_raw": raw.strip(), "name": m.group(1), "price": m.group(2),
                                 "desc": None, "desc_raw": None, "gap": gap})
            blanks = 0
            continue
        if "(" in t or ")" in t:
            it = page["sections"][-1]["items"][-1]
            it["desc_raw"], it["desc"] = t, t
            blanks = 0
            continue
        raise ValueError(f"neprepoznana vrstica: {t!r}")
    return bars


def clean(bars):
    """Uporabi popravke in zabeleži spremembe za poročilo."""
    log = []
    for b in BARS:
        data = bars[b["docx"]]
        for pi, page in enumerate(data["pages"], 1):
            for sec in page["sections"]:
                for it in sec["items"]:
                    price_fmt = it["price"] + NBSP + "€"
                    orig_price_txt = re.search(r"(\d+,\d\s*€)\s*$", it["name_raw"]).group(1)
                    it["price_txt"] = price_fmt
                    new_name, why_n = fix_text(it["name"])
                    it["name_fixed"] = new_name
                    spacing = re.search(r"[^\s\d](\d+,\d)", it["name_raw"])
                    if spacing:
                        why_n.append("manjkajoč presledek pred ceno")
                    if why_n:
                        log.append((b["name"], pi, it["name_raw"], new_name + " " + price_fmt.replace(NBSP, " "), why_n))
                    if it["desc"]:
                        new_desc, why_d = fix_text(it["desc"])
                        it["desc_fixed"] = new_desc
                        if why_d:
                            log.append((b["name"], pi, it["desc_raw"], new_desc, why_d))
                    else:
                        it["desc_fixed"] = None
    return log


# ---- Logotipi --------------------------------------------------------------
def load_logos(svg_path):
    root = ET.parse(svg_path).getroot()
    css = root.find(f"{{{NS}}}defs/{{{NS}}}style").text
    cls = {}
    for sel, body in re.findall(r"([^{}]+)\{([^{}]*)\}", css):
        props = dict(re.findall(r"([\w-]+)\s*:\s*([^;]+);?", body))
        for s in re.findall(r"\.(st\d+)", sel):
            cls.setdefault(s, {}).update({k: v.strip() for k, v in props.items()})

    def inline(el, recolor=None):
        for e in el.iter():
            c = e.attrib.pop("class", None)
            if c:
                p = cls.get(c, {})
                fill = p.get("fill", "#000000")
                if recolor and fill.lower() in ("#231f20",):
                    fill = recolor
                e.set("fill", fill)
                if p.get("fill-rule"):
                    e.set("fill-rule", p["fill-rule"])
        return el

    def ser(el):
        s = ET.tostring(el, encoding="unicode")
        return s.replace(f' xmlns="{NS}"', "").replace(f"ns0:", "")

    groups = {g.get("id"): g for g in root if g.tag == f"{{{NS}}}g"}

    def wrap(children, gid, recolor):
        out = []
        for c in children:
            out.append(ser(inline(c, recolor)))
        return f'<g id="{gid}">' + "".join(out) + "</g>"

    logos = {}
    for key in ("MARS", "LUNA"):
        kids = list(groups[key])
        logos[key] = dict(svg=wrap(kids, f"logotip_{key.lower()}", WHITE),
                          box=(0 if key == "MARS" else 1092.43, 105.9,
                               819.7 if key == "MARS" else 705.28, 349.0))
    kids = list(groups["SONCE"])
    logos["BALTHAZAR"] = dict(svg=wrap([kids[4]], "logotip_balthazar", None), box=(3175, 0, 879, 609))
    return logos


def place_logo(logo, cx, top, width):
    x, y, w, h = logo["box"]
    s = width / w
    px = cx - width / 2
    return f'<g transform="translate({px - x * s:.2f} {top - y * s:.2f}) scale({s:.5f})">{logo["svg"]}</g>', h * s


# ---- Dekor: zima -----------------------------------------------------------
def flake(cx, cy, r, rot, color, op, sw):
    d = []
    for k in range(6):
        a = math.radians(rot + 60 * k)
        ux, uy = math.cos(a), math.sin(a)
        d.append(f"M{cx:.2f} {cy:.2f}L{cx + ux * r:.2f} {cy + uy * r:.2f}")
        for frac, ln in ((0.58, 0.30), (0.82, 0.18)):
            bx, by = cx + ux * r * frac, cy + uy * r * frac
            for sgn in (-1, 1):
                b = a + sgn * math.radians(60)
                d.append(f"M{bx:.2f} {by:.2f}L{bx + math.cos(b) * r * ln:.2f} {by + math.sin(b) * r * ln:.2f}")
    return (f'<path d="{"".join(d)}" fill="none" stroke="{color}" stroke-opacity="{op}" '
            f'stroke-width="{sw}" stroke-linecap="round"/>')


def decor(bar, pno, seed):
    rnd = random.Random(seed)
    out = []
    # mehek zimski sij
    out.append(f'<rect id="sij" width="{W}" height="{H}" fill="url(#sij_{bar["id"]}{pno})"/>')
    flakes = []
    big = [(W - 26, 52, 70), (-4, 112, 40)] if pno == 1 else [(26, 52, 62), (W + 2, 128, 38)]
    for i, (x, y, r) in enumerate(big):
        flakes.append(flake(x, y, r, rnd.uniform(0, 60), TEAL if i % 2 == 0 else WHITE, 0.30 if i % 2 == 0 else 0.16, 0.9))
    for i in range(9):
        side = rnd.choice((0, 1))
        x = rnd.uniform(8, 26) if side == 0 else rnd.uniform(W - 26, W - 8)
        y = rnd.uniform(110, H - 90)
        flakes.append(flake(x, y, rnd.uniform(4, 10), rnd.uniform(0, 60), WHITE, rnd.choice((0.18, 0.28, 0.4)), 0.6))
    for i in range(18):
        x = rnd.uniform(6, W - 6)
        y = rnd.uniform(6, 120) if rnd.random() < 0.6 else rnd.uniform(6, H - 6)
        if M - 4 < x < W - M + 4 and y > 150:
            x = rnd.choice((rnd.uniform(6, M - 8), rnd.uniform(W - M + 8, W - 6)))
        flakes.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{rnd.choice((0.7, 1.0, 1.4, 1.8))}" fill="{WHITE}" fill-opacity="{rnd.choice((0.2, 0.35, 0.5))}"/>')
    out.append('<g id="snezinke">' + "".join(flakes) + "</g>")
    # snežni zamet
    y0 = H - 78
    d1 = f"M0 {y0 + 8}C{W * .2:.1f} {y0 - 8} {W * .38:.1f} {y0 + 14} {W * .6:.1f} {y0 + 2}S{W * .92:.1f} {y0 - 10} {W} {y0 + 4}V{H}H0Z"
    out.append(f'<path id="zamet" d="{d1}" fill="{TEAL}" fill-opacity="0.10"/>')
    return "".join(out)


# ---- Besedilo --------------------------------------------------------------
def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def text(x, y, s, size, weight=400, fill=WHITE, anchor="start", op=None, ls=None, gid=None):
    fam, w = FONT[weight]
    a = f' text-anchor="{anchor}"' if anchor != "start" else ""
    o = f' fill-opacity="{op}"' if op is not None else ""
    l = f' letter-spacing="{ls}"' if ls is not None else ""
    i = f' id="{gid}"' if gid else ""
    return (f'<text{i} x="{x:.2f}" y="{y:.2f}" font-family="{fam}, Archivo, sans-serif" font-weight="{w}" '
            f'font-size="{size}" fill="{fill}"{o}{a}{l}>{esc(s)}</text>')


def layout_page(bar, page, pno, logo, seed):
    n = bar["id"]
    parts = []
    # ozadje
    parts.append(f'<rect id="ozadje" width="{W}" height="{H}" fill="{NAVY}"/>')
    parts.append(f'<g id="zimski_dekor">{decor(bar, pno, seed)}</g>')
    # logotip
    top, cx = 36.0, W / 2
    if bar["logo"] == "BALTHAZAR":
        r = 76
        parts.append(f'<circle id="luna_plosca" cx="{cx}" cy="{top + r}" r="{r}" fill="#f7f5f5"/>')
        svg, lh = place_logo(logo, cx, top + r - (609 * 120 / 879) / 2, 120)
        parts.append(svg)
        logo_bottom = top + 2 * r
    else:
        svg, lh = place_logo(logo, cx, top, 170)
        parts.append(svg)
        logo_bottom = top + lh
    # vsebina
    zone_top = logo_bottom + 30
    zone_bot = H - 62 - 14
    secs = page["sections"]
    n_items = sum(len(s["items"]) for s in secs)
    n_heads = sum(1 for s in secs if s["title"])
    n_gaps = sum(1 for s in secs for it in s["items"] if it["gap"])
    head_h, sec_gap, item_gap = 26.0, 20.0, 12.0
    fixed = head_h * n_heads + sec_gap * max(0, len(secs) - 1) + item_gap * n_gaps
    pitch = max(18.0, min(38.0, (zone_bot - zone_top - fixed) / n_items))
    ns, ds = (9.6, 7.0) if pitch < 24 else (10.5, 7.8)
    y = zone_top  # zgornji rob naslednjega bloka
    body = []
    for si, s in enumerate(secs):
        if si:
            y += sec_gap
        if s["title"]:
            body.append(text(M, y + 10, s["title"], 11.5, 600, bar["head"], gid=f"naslov_{re.sub(r'[^a-z]', '', s['title'].lower())}"))
            body.append(f'<line x1="{M}" y1="{y + 16:.2f}" x2="{W - M}" y2="{y + 16:.2f}" stroke="{bar["head"]}" stroke-opacity="0.55" stroke-width="0.6"/>')
            y += head_h
        for it in s["items"]:
            if it["gap"]:
                y += item_gap
            body.append(text(M, y + ns, it["name_fixed"], ns, 600))
            body.append(text(W - M, y + ns, it["price_txt"], ns, 700, bar["price"], anchor="end"))
            if it["desc_fixed"]:
                body.append(text(M, y + ns + ds + 3.2, it["desc_fixed"], ds, 400, op=0.66))
            y += pitch
    free = zone_bot - y
    if free > 110:
        r = min(72.0, free / 2 - 12)
        parts.append(f'<g id="snezinka_velika">{flake(W / 2, (y + zone_bot) / 2, r, 15, TEAL, 0.22, 0.9)}</g>')
    parts.append(f'<g id="{n}_s{pno}_vsebina">' + "".join(body) + "</g>")
    # podnožje
    fy = H - 62
    foot = [f'<line x1="{M}" y1="{fy}" x2="{W - M}" y2="{fy}" stroke="{WHITE}" stroke-opacity="0.28" stroke-width="0.5"/>',
            text(M, fy + 15, SLOGAN, 8.5, 600, TEAL, ls=0.4, gid="slogan"),
            text(M, fy + 30, FOOT[0], 6.3, 400, op=0.7, gid="podnozje_1"),
            text(M, fy + 38.5, FOOT[1], 6.3, 400, op=0.7, gid="podnozje_2")]
    parts.append(f'<g id="{n}_s{pno}_podnozje">' + "".join(foot) + "</g>")
    return "".join(parts)


def build_svg(bars, logos):
    TW, TH = 2 * W + GAP, 3 * H + 2 * GAP
    defs = []
    for b in BARS:
        for pno in (1, 2):
            gx = W * (0.5 if pno == 1 else 0.35)
            defs.append(f'<radialGradient id="sij_{b["id"]}{pno}" gradientUnits="userSpaceOnUse" cx="{gx:.1f}" cy="0" r="{H * 0.62:.1f}">'
                        f'<stop offset="0" stop-color="{TEAL}" stop-opacity="0.34"/>'
                        f'<stop offset="1" stop-color="{TEAL}" stop-opacity="0"/></radialGradient>')
    out = [f'<?xml version="1.0" encoding="UTF-8"?>\n<svg xmlns="{NS}" viewBox="0 0 {TW:.2f} {TH:.2f}" width="{TW:.2f}" height="{TH:.2f}">',
           "<title>Zimski letaki Odiseja: LUNA, MARS, Balthazar (148 × 220 mm)</title>",
           "<defs>" + "".join(defs) + "</defs>"]
    seed = 7
    for ri, b in enumerate(BARS):
        for pno, page in enumerate(bars[b["docx"]]["pages"], 1):
            ox, oy = (pno - 1) * (W + GAP), ri * (H + GAP)
            seed += 1
            out.append(f'<g id="{b["id"]}_stran{pno}" transform="translate({ox:.2f} {oy:.2f})">'
                       + layout_page(b, page, pno, logos[b["logo"]], seed) + "</g>")
    out.append("</svg>")
    return "\n".join(out)


# ---- Excel in poročilo -----------------------------------------------------
def build_xlsx(bars, path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for b in BARS:
        ws = wb.create_sheet(b["name"])
        ws.append(["Produkt", "Vsebina", "Cena"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="071D22")
        for page in bars[b["docx"]]["pages"]:
            for s in page["sections"]:
                if s["title"]:
                    ws.append([s["title"], None, None])
                    for c in ws[ws.max_row]:
                        c.font = Font(bold=True, color="00868F")
                        c.fill = PatternFill("solid", fgColor="E6F6F7")
                for it in s["items"]:
                    desc = it["desc_fixed"]
                    if desc:
                        desc = desc[1:-1] if desc.startswith("(") and desc.endswith(")") else desc
                    ws.append([it["name_fixed"], desc, float(it["price"].replace(",", "."))])
                    ws.cell(ws.max_row, 3).number_format = '0.00" €"'
                    ws.cell(ws.max_row, 3).alignment = Alignment(horizontal="right")
        ws.freeze_panes = "A2"
        for col, wd in zip("ABC", (46, 72, 10)):
            ws.column_dimensions[col].width = wd
    wb.save(path)


def build_report(log, bars, path):
    L = ["# Poročilo: zimski letaki LUNA, MARS, Balthazar", "",
         "Besedilo je prepisano iz `Zimski_letak_bari.docx`. Popravljene so samo slovnične in tipkarske napake; "
         "sestave pijač, naslovi razdelkov in vrstni red so ostali kot v dokumentu.", "",
         "## Popravki besedila", "", "| Lokal | Stran | Prej | Potem | Razlog |", "|---|---|---|---|---|"]
    for bar, pno, before, after, why in log:
        L.append(f"| {bar} | {pno} | {before} | {after} | {'; '.join(dict.fromkeys(why))} |")
    L += ["", "### Zapis cen", "",
          "Vse cene so zapisane enotno kot »4,8 €« (nedeljiv presledek med številko in znakom €). "
          "Število decimalk je ostalo kot v dokumentu (npr. »7,0 €«, »10,0 €«).", "",
          "## Prepisano brez sprememb (vsebinske razlike med lokali)", "",
          "- Naslov prvega razdelka: LUNA »Topli brezalkoholni napitki«; MARS in Balthazar »Topli napitki«.",
          "- Screwballada: LUNA »vroča čokolada, whiskey Skrewball, smetana, arašidi«; MARS »vroča čokolada, smetana, arašidi, Skrewball«. "
          "Sestava je ostala različna, kot v dokumentu.",
          "- Vroča limonada z ingverjem in medom je pri LUNA brez razdelka (stoji za vročimi čokoladami, ločena z odmikom); "
          "pri MARS in Balthazar je med toplimi napitki.",
          "- MARS: alkoholni topli napitki (Kuhan gin, Vroči punč, Kuhano vino) so v istem razdelku kot brezalkoholni, kot v dokumentu.",
          "- Zapis imen koktejlov (»Moscow mule«, »Porn star martini«, »Tom Collins«) je ostal kot v dokumentu.", "",
          "## Odločitve pri izdelavi", "",
          "- Lokali: LUNA, MARS, Balthazar. SONCE je v datoteki z logotipi, a v dokumentu nima vsebine, zato ni letaka.",
          "- V dokumentu piše »Baltazar«, na logotipu »Balthazar«. Na letaku je logotip; ime ni zapisano kot besedilo.",
          "- Blagovna znamka: »Screwball« in »screwball« sta poenotena v »Skrewball« (pravilen zapis znamke, kot pri »Skrewball espresso martini«). "
          "Ime pijače »Screwballada« je ostalo.",
          "- »radenska« je zapisana kot »Radenska« (blagovna znamka).",
          "- Besedilo podnožja je na vseh 6 straneh. Slogan »Vesolje vsebin« je na vsaki strani nad podnožjem.",
          "- Logotip Balthazar ima črne dele, ki bi na temno modrem ozadju izginili, zato je na svetli plošči (#f7f5f5, barva iz datoteke z logotipi). "
          "Logotipa LUNA in MARS sta bela.",
          "- Barve CGP Odiseje: navy #071d22, teal #00b1bc, zelena #a5cb39; ostalo so bela in prosojnosti.", "",
          "## Tehnično", "",
          "- SVG ima vseh 6 strani (148 × 220 mm = 419,53 × 623,62 pt) ena ob drugi: vrstice LUNA, MARS, Balthazar; levo prva stran, desno druga. "
          "Razmik med stranmi je 20 mm. Enote so točke, 1 uporabniška enota = 1 pt.",
          "- Vsak lokal in stran je svoja skupina (`luna_stran1` …), znotraj nje ozadje, zimski dekor, logotip, vsebina in podnožje. "
          "Besedilo je živo (`<text>`).",
          "- Pisava je v datoteki navedena kot Archivo (Regular, Medium, SemiBold, Bold) in ni vdelana. "
          "Če jo Illustrator ob odprtju zamenja, jo zamenjajte prek Type > Find Font.",
          "- Za Illustrator ustvarite artboarde po obstoječih ozadjih strani (Object > Artboards > Convert to Artboards po izbiri ozadij).",
          ""]
    Path(path).write_text("\n".join(L), encoding="utf8")


def main():
    docx, logo_svg = sys.argv[1], sys.argv[2]
    out = Path(sys.argv[3] if len(sys.argv) > 3 else "letaki")
    out.mkdir(parents=True, exist_ok=True)
    bars = read_docx(docx)
    log = clean(bars)
    logos = load_logos(logo_svg)
    (out / "zimski_letaki.svg").write_text(build_svg(bars, logos), encoding="utf8")
    build_xlsx(bars, out / "zimska_ponudba.xlsx")
    build_report(log, bars, out / "porocilo.md")
    for b in BARS:
        print(b["name"], [sum(len(s["items"]) for s in p["sections"]) for p in bars[b["docx"]]["pages"]])
    print(len(log), "popravkov")


if __name__ == "__main__":
    main()
